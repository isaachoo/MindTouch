#!/usr/bin/env python3
"""
Compare two carers.hk exports (from scrape_carers.py) and write a change report.

  python3 scripts/diff_carers.py exports/carers_2026-09-21.csv exports/carers_2026-09-28.csv -o reports/changes_2026-09-28

Writes <out>.xlsx (sheets: Summary, Added, Removed, Changed, Current) when openpyxl is installed,
otherwise <out>_added.csv, <out>_removed.csv, <out>_changed.csv. Units are matched by unit_id.
"""
from __future__ import annotations

import argparse
import csv
import sys
from datetime import date

KEY = "unit_id"
# Fields worth reporting when they change. detail_text is compared separately (too long to show).
WATCH = ["name", "address", "district", "tel", "website", "tags", "opening_time", "categories", "subcategories", "detail_url"]
SHOW = ["unit_id", "name", "district", "address", "tel", "website", "categories", "subcategories", "detail_url"]


def load(path: str) -> dict[str, dict]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    if rows and KEY not in rows[0]:
        sys.exit(f"{path}: no {KEY} column; is this a scrape_carers.py export?")
    return {r[KEY]: r for r in rows}


def diff(old: dict[str, dict], new: dict[str, dict]):
    added = [new[k] for k in new if k not in old]
    removed = [old[k] for k in old if k not in new]
    changed = []
    for k in new:
        if k not in old:
            continue
        o, n = old[k], new[k]
        for field in WATCH:
            if (o.get(field) or "") != (n.get(field) or ""):
                changed.append({"unit_id": k, "name": n.get("name", ""), "field": field, "old": o.get(field, ""), "new": n.get(field, ""), "detail_url": n.get("detail_url", "")})
        if (o.get("detail_text") or "") != (n.get("detail_text") or ""):
            changed.append({"unit_id": k, "name": n.get("name", ""), "field": "detail_text", "old": "(page text changed)", "new": f"{len(n.get('detail_text') or '')} chars", "detail_url": n.get("detail_url", "")})
    return added, removed, changed


def write_csvs(out: str, added, removed, changed):
    for name, rows, cols in [("added", added, SHOW), ("removed", removed, SHOW), ("changed", changed, ["unit_id", "name", "field", "old", "new", "detail_url"])]:
        with open(f"{out}_{name}.csv", "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=cols, quoting=csv.QUOTE_ALL, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
    print(f"wrote {out}_added.csv, {out}_removed.csv, {out}_changed.csv", file=sys.stderr)


def write_xlsx(out: str, added, removed, changed, current: list[dict], summary: list[tuple[str, str]]):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for r in summary:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60

    def sheet(title, rows, cols, fill=None):
        s = wb.create_sheet(title)
        s.append(cols)
        for c in s[1]:
            c.font = Font(bold=True)
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)
        for r in rows:
            s.append([(r.get(c) or "")[:32000] for c in cols])
        for i, c in enumerate(cols, 1):
            s.column_dimensions[get_column_letter(i)].width = 14 if c in ("unit_id", "tel", "district", "field") else 40
        s.freeze_panes = "A2"
        s.auto_filter.ref = s.dimensions

    sheet("Added", added, SHOW, "C6EFCE")
    sheet("Removed", removed, SHOW, "FFC7CE")
    sheet("Changed", changed, ["unit_id", "name", "field", "old", "new", "detail_url"], "FFEB9C")
    current_cols = [c for c in (current[0].keys() if current else SHOW) if c != "detail_text"]
    sheet("Current", current, current_cols)
    path = f"{out}.xlsx"
    wb.save(path)
    print(f"wrote {path}", file=sys.stderr)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("old")
    ap.add_argument("new")
    ap.add_argument("-o", "--out", default=f"changes_{date.today().isoformat()}", help="output path without extension")
    args = ap.parse_args()

    old, new = load(args.old), load(args.new)
    added, removed, changed = diff(old, new)
    summary = [
        ("Report", args.out),
        ("Previous export", f"{args.old} ({len(old)} units)"),
        ("Current export", f"{args.new} ({len(new)} units)"),
        ("Added units", str(len(added))),
        ("Removed units", str(len(removed))),
        ("Changed fields", str(len(changed))),
        ("Units with changes", str(len({c['unit_id'] for c in changed}))),
        ("Source", "照顧者資訊網 carers.hk"),
    ]
    for k, v in summary:
        print(f"{k:18} {v}", file=sys.stderr)

    try:
        write_xlsx(args.out, added, removed, changed, list(new.values()), summary)
    except ImportError:
        print("openpyxl not installed (pip install openpyxl); writing CSVs instead", file=sys.stderr)
        write_csvs(args.out, added, removed, changed)
    return 0


if __name__ == "__main__":
    sys.exit(main())
